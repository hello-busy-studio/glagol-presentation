import sys, copy
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
import os
LOGO_WHITE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'assets', 'logo', 'logo-mono-white.png')
prs = Presentation(SRC)
ORDER = [21, 46, 43, 75, 16, 67, 78, 59, 81, 84, 69, 47, 92]

# --- keep only chosen slides, in order; drop template comments
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
keep = [ids[n - 1] for n in ORDER]
for el in ids:
    sldIdLst.remove(el)
    if el not in keep:
        prs.part.drop_rel(el.get(qn('r:id')))
for el in keep:
    sldIdLst.append(el)
for sl in prs.slides:
    for rId, rel in list(sl.part.rels.items()):
        if rel.reltype.endswith('/comments'):
            sl.part.drop_rel(rId)

def find(slide, sid):
    def walk(shs):
        for s in shs:
            if s.shape_id == sid: return s
            if s.shape_type == 6:
                r = walk(s.shapes)
                if r is not None: return r
    s = walk(slide.shapes)
    assert s is not None, sid
    return s

def set_text(shape, paras, accent=None):
    """paras: list of paragraphs; each paragraph is a str or list of lines (joined by line breaks).
    accent: set of (para_idx, line_idx) to paint accent1. Formatting is cloned from the first run."""
    txBody = shape.text_frame._txBody
    ps = txBody.findall(qn('a:p'))
    tpl_p = ps[0]
    tpl_r = tpl_p.find(qn('a:r'))
    tpl_rPr = tpl_r.find(qn('a:rPr'))
    for p in ps: txBody.remove(p)
    for pi, para in enumerate(paras):
        lines = [para] if isinstance(para, str) else para
        p = copy.deepcopy(tpl_p)
        for ch in list(p):
            if ch.tag in (qn('a:r'), qn('a:br'), qn('a:fld')): p.remove(ch)
        end = p.find(qn('a:endParaRPr'))
        for li, line in enumerate(lines):
            if li:
                br = p.makeelement(qn('a:br'), {}); br.append(copy.deepcopy(tpl_rPr))
                end.addprevious(br) if end is not None else p.append(br)
            r = copy.deepcopy(tpl_r)
            r.find(qn('a:t')).text = line
            rPr = r.find(qn('a:rPr'))
            if rPr is not None and rPr.get('u'): del rPr.attrib['u']
            if accent and (pi, li) in accent:
                sf = rPr.find(qn('a:solidFill'))
                for c in list(sf): sf.remove(c)
                sc = sf.makeelement(qn('a:schemeClr'), {'val': 'accent1'}); sf.append(sc)
            end.addprevious(r) if end is not None else p.append(r)
        txBody.append(p)

def placeholder(slide, sid, label, round_=False, white=False):
    """Replace a picture with a visible grey placeholder that names what the client must send."""
    pic = find(slide, sid)
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
                                 pic.left, pic.top, pic.width, pic.height)
    if round_: shp.adjustments[0] = 0.16667
    shp.fill.solid()
    if white: shp.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    else: shp.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_5
    shp.line.fill.background(); shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.name = 'Wix Madefor Display'; r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x72, 0x72, 0x72)
    pic._element.addprevious(shp._element)          # keep z-order
    pic._element.getparent().remove(pic._element)

def style_tag(shape, accent=False):
    """Tag per DESIGN.md §5.1: pill h=0.24in, 0.75pt accent1 outline, 8pt text; accent=True -> filled accent1, text lt1."""
    shape.height = Emu(int(0.24 * 914400))
    shape.fill.solid() if accent else shape.fill.background()
    if accent: shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    shape.line.width = Pt(0.75); shape.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(8)
            r.font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1 if accent else MSO_THEME_COLOR.ACCENT_1

def delete(slide, sid):
    el = find(slide, sid)._element; el.getparent().remove(el)

s = prs.slides
# 1 — обложка (шаблон 21, рыжая)
sl = s[0]
set_text(find(sl, 271), [['Как защитить', 'бюджет', 'на обучение']])
set_text(find(sl, 272), [['Онлайн-эфир бюро Глагол', '8 октября, 19:00', 'Спикер — Лера Вишневская']])
placeholder(sl, 273, '[фото к обложке]')
# on orange the base logo's orange arrow disappears -> mono white logo
for sh in list(sl.shapes):
    if sh.shape_type == 13 and sh.height < Emu(int(0.25*914400)) and sh.top > Emu(int(4.5*914400)) and sh.left < Emu(914400):
        new = sl.shapes.add_picture(LOGO_WHITE, sh.left, sh.top, height=sh.height)
        sh._element.addprevious(new._element); sh._element.getparent().remove(sh._element)

# 2 — спикер (46)
sl = s[1]
placeholder(sl, 566, '[фото спикера]')
placeholder(sl, 567, '[QR-код]', round_=True, white=True)
set_text(find(sl, 569), ['спикер'])
set_text(find(sl, 570), ['Яндекс'])
set_text(find(sl, 571), ['soft skills'])
set_text(find(sl, 572), ['Лера Вишневская']); find(sl, 572).width = Emu(int(4.0*914400))
set_text(find(sl, 573), [['Руководитель образовательных продуктов soft skills в Яндексе']])
r573 = find(sl, 573); r573.width = Emu(int(3.82*914400))
for r in r573.text_frame.paragraphs[0].runs: r.font.size = Pt(12)
# tags: outline + one accent (DESIGN.md §5.1), widths fit text, gap 0.1in
x = Emu(int(0.543*914400))
for sid, w, acc in ((570, 0.723, False), (571, 0.835, True)):
    tg = find(sl, sid); style_tag(tg, accent=acc); tg.left = x; tg.top = Emu(int(1.029*914400)); tg.width = Emu(int(w*914400))
    x = tg.left + tg.width + Emu(int(0.1*914400))

# 3 — план эфира (43)
sl = s[2]
set_text(find(sl, 541), ['план эфира'])
for sid, lines in zip((531, 534, 537, 540), (['Разбор', 'на примерах'], ['Возражения', 'бизнеса'],
                                              ['Чек-лист', 'критериев'], ['Ответы', 'на вопросы'])):
    set_text(find(sl, sid), [lines])

# 4 — кому полезно (75)
sl = s[3]
set_text(find(sl, 875), ['аудитория'])
set_text(find(sl, 857), [['Кому', 'будет', 'полезно']])
set_text(find(sl, 858), [['Эфир для тех, кому предстоит', 'защищать бюджет на обучение', 'на 2027 год']])
items = [['Руководителям T&D, L&D и HR, у которых', 'впереди защита бюджета на обучение'],
         ['Руководителям, которые защищают', 'бюджет своего направления'],
         ['Тем, кто помогает командам готовить', 'защиты проектов и бюджетов']]
for sid, lines in zip((861, 864, 867), items):
    set_text(find(sl, sid), [lines])
delete(sl, 869); delete(sl, 872)

# 5 — разделитель (16)
sl = s[4]
set_text(find(sl, 227), [['Разбор', 'на примерах']])
set_text(find(sl, 228), [['[Короткое описание блока:', 'о чем пойдет речь, 1–2 строки]']]); find(sl, 228).top = Emu(int(1.85*914400))

# 6 — кейс: текст + фото (67)
sl = s[5]
set_text(find(sl, 766), ['кейс'])
set_text(find(sl, 763), [['[Название', 'кейса]']])
set_text(find(sl, 764), [['[Компания, задача и сумма', 'бюджета. Как проходила защита', 'и чем закончилась — 3–4 строки]']])
placeholder(sl, 762, '[фото или скрин кейса]')

# 7 — критерии: 3 карточки (78)
sl = s[6]
set_text(find(sl, 939), ['критерии защиты'])
set_text(find(sl, 925), [['Критерии', 'эффективной', 'защиты']])
set_text(find(sl, 926), [['Применимы к любой защите —', 'соберем их в чек-лист', 'по ходу эфира']])
for i, (t, d) in enumerate(((928, 929), (932, 933), (936, 937)), 1):
    set_text(find(sl, t), [f'[Критерий {i}]'])
    set_text(find(sl, d), ['[Пояснение', 'в одну-две строки]'])

# 8 — важная мысль (59)
sl = s[7]
set_text(find(sl, 690), ['главная мысль'])
set_text(find(sl, 689), [['[Ключевая мысль', 'в две строки]']])

# 9 — возражения → ответы (81)
sl = s[8]
set_text(find(sl, 1016), ['возражения бизнеса в этом сезоне'])
for i, (c, d) in enumerate(((1000, 1002), (1004, 1006), (1008, 1010), (1012, 1014)), 1):
    set_text(find(sl, c), [f'[Возражение {i}', 'в одну-две строки]'])
    set_text(find(sl, d), ['[Чем отвечать: аргумент', 'в одну-две строки]'])

# 10 — таймлайн подготовки (84)
sl = s[9]
set_text(find(sl, 1076), ['как готовиться к защите'])
for i, (c, d, dt) in enumerate(((1060, 1062, 1063), (1064, 1073, 1070), (1066, 1074, 1071), (1068, 1075, 1072)), 1):
    set_text(find(sl, c), [f'[Этап {i}', 'название этапа]'])
    set_text(find(sl, d), ['[Что сделать на этом этапе:', 'цифры, аргументы,', 'с кем согласовать]'])
    set_text(find(sl, dt), ['[срок]'])

# 11 — цитата (69)
sl = s[10]
set_text(find(sl, 785), ['цитата'])
set_text(find(sl, 783), [['[Цитата спикера', 'в две или', 'три строки]']])
set_text(find(sl, 784), [['— Лера Вишневская']])

# 12 — о бюро (47): реальные данные из шаблона, без изменений, кроме метки
set_text(find(s[11], 588), ['о бюро Глагол'])

# 13 — финал (92)
sl = s[12]
set_text(find(sl, 1183), [['Спасибо', 'за внимание!']])
set_text(find(sl, 1184), [['Заберите чек-лист критериев', 'и презентацию с эфира']])
placeholder(sl, 1188, '[QR на материалы]')
set_text(find(sl, 1189), ['[ссылка]'])

prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
